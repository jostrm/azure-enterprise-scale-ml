"""Deterministic local fixture generation, with a separate Decimal evaluation oracle.

Workbook normalization reads the saved workbook, never its XML representation.
The image is an original geometric drawing. Its DOCX wrapper is minimal OOXML
written with the standard library; no Office installation or model call is used.
"""

from datetime import date, datetime, timedelta
from decimal import Decimal, localcontext
import hashlib
from io import BytesIO
import json
from pathlib import Path
import random
import re
import shutil
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import openpyxl
from openpyxl.styles import Font
from PIL import Image, ImageDraw
import tiktoken

from .arithmetic import IDENTIFIER, validate_row


DATASET = "car-sales"
FIELDS = (
    "dataset", "version", "sale_id", "dealership", "date", "car_model", "quantity",
    "unit_price_minor", "discount_minor", "refund_minor", "tax_minor", "status", "currency",
)
INPUT_PATHS = {
    "json": "ingestion/json/retail/sales-json/car-sales-xyz.json",
    "excel": "ingestion/excel/retail/sales-excel/car-calculator-tool.txt",
}
LIBRARY_LAYOUT = (
    "ingestion/json/retail/sales-json",
    "ingestion/excel/retail/sales-excel",
    "ingestion/json/retail/service-notices",
    "ingestion/json/corporate/vehicle-catalog",
    "ingestion/excel/retail/finance-guides",
    "ingestion/excel/corporate/fleet-policy",
)
IMAGE_FACT = "A blue car has a magenta triangle marked 73 on its door."
FIXED_TIMESTAMP = "2000-01-01T00:00:00Z"


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json(value, *, indent=2) -> str:
    return json.dumps(value, ensure_ascii=False, indent=indent, allow_nan=False) + "\n"


def _write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def _sale(rng: random.Random, number: int, version: str, *, fixture=False) -> dict:
    quantity = rng.randint(1, 5)
    price = rng.randint(1800000, 9900000)
    gross = quantity * price
    discount = rng.randint(0, gross // 12)
    status = "cancelled" if number % 19 == 0 else "refunded" if number % 11 == 0 else "completed"
    refund = rng.randint(1, gross - discount) if status == "refunded" else 0
    row = {
        "dataset": DATASET, "version": version,
        "sale_id": f"{'FIXTURE' if fixture else 'SALE'}-{number:06d}",
        "dealership": rng.choice(("North Motors", "Riverside Autos", "Summit Vehicles", "Cedar Garage")),
        "date": (date(2026, 1, 1) + timedelta(days=rng.randrange(270))).isoformat(),
        "car_model": rng.choice(("Cedar EV", "Orion Hybrid", "Harbor Touring", "Summit Utility",
                                 "Aspen Compact", "Café Élan")),
        "quantity": quantity, "unit_price_minor": price, "discount_minor": discount,
        "refund_minor": refund, "tax_minor": (gross - discount - refund) * 7 // 100,
        "status": status, "currency": "USD",
    }
    validate_row(row, DATASET, version)
    return row


def _json_table(rows: list[dict], version: str) -> str:
    return _json({"dataset": DATASET, "version": version, "currency": "USD", "records": rows})


def _worksheet_text(rows: list[dict]) -> str:
    lines = ["worksheet=Car sales", "schema=car-sales-normalized-v1"]
    for index, row in enumerate(rows, 1):
        lines.append(f"ROW {index}")
        lines.extend(f"{field}={json.dumps(row[field], ensure_ascii=False)}" for field in FIELDS)
        lines.append("END ROW")
    return "\n".join(lines) + "\n"


def _partition_rows(source: str, version: str, seed: int, encoding) -> list[dict]:
    # Separate RNG streams prevent changing the JSON row count from reshuffling Excel.
    rng = random.Random(f"{DATASET}:{version}:{seed}:{source}")
    rows = []
    number_base = 0 if source == "json" else 100000
    for index in range(1, 1001):
        rows.append(_sale(rng, number_base + index, version))
        text = _json_table(rows, version) if source == "json" else _worksheet_text(rows)
        count = len(encoding.encode(text, disallowed_special=()))
        if 13500 <= count <= 14500:
            return rows
        if count > 14500:
            raise ValueError("a complete additional row overshot the deterministic input token budget")
    raise ValueError("failed to produce a sufficiently large complete table")


def _stable_zip(entries: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(entries):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0o600 << 16
            archive.writestr(info, entries[name], compress_type=ZIP_DEFLATED, compresslevel=9)
    return buffer.getvalue()


def _save_workbook(path: Path, rows: list[dict]) -> list[dict]:
    workbook = openpyxl.Workbook()
    workbook.properties.creator = "Synthetic sales fixture"
    workbook.properties.lastModifiedBy = "Synthetic sales fixture"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    sheet = workbook.active
    sheet.title = "Car sales"
    sheet.append(FIELDS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row[field] for field in FIELDS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    with ZipFile(BytesIO(buffer.getvalue())) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    core = entries["docProps/core.xml"].decode("utf-8")
    core = re.sub(r"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                  lambda match: match[1] + FIXED_TIMESTAMP + match[2], core)
    entries["docProps/core.xml"] = core.encode("utf-8")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_stable_zip(entries))
    saved = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        iterator = saved["Car sales"].iter_rows(values_only=True)
        header = tuple(next(iterator))
        if header != FIELDS:
            raise ValueError("saved workbook schema does not match the authoritative schema")
        actual = [dict(zip(header, values)) for values in iterator]
        for row in actual:
            validate_row(row, DATASET, rows[0]["version"])
        if actual != rows:
            raise ValueError("saved workbook rows differ from source records")
        return actual
    finally:
        saved.close()


def _draw_car() -> bytes:
    image = Image.new("RGB", (900, 480), "#edf5fc")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 370, 900, 480), fill="#c5cdd5")
    draw.line((0, 417, 900, 417), fill="white", width=7)
    draw.polygon([(130, 285), (240, 255), (325, 150), (590, 150),
                  (690, 255), (790, 280), (810, 355), (105, 355)], fill="#1879cf", outline="#113c68")
    draw.polygon([(280, 250), (342, 170), (445, 170), (445, 250)], fill="#b5e8f7")
    draw.polygon([(465, 170), (577, 170), (652, 250), (465, 250)], fill="#b5e8f7")
    draw.line((455, 165, 455, 350), fill="#113c68", width=5)
    draw.rounded_rectangle((380, 266, 421, 274), radius=3, fill="#163c58")
    for center in (250, 685):
        draw.ellipse((center - 60, 300, center + 60, 420), fill="#20232a")
        draw.ellipse((center - 34, 326, center + 34, 394), fill="#bbc6cf")
        draw.ellipse((center - 11, 349, center + 11, 371), fill="#4c5964")
    draw.rectangle((764, 284, 796, 307), fill="#fff0ac")
    draw.rectangle((112, 286, 139, 310), fill="#db2929")
    draw.polygon([(540, 235), (465, 350), (615, 350)], fill="#dd168d", outline="#6e0848")
    # Hand-drawn seven-segment numerals avoid platform fonts and embedded answer text.
    segments = {
        "a": ((0, 0), (22, 0)), "b": ((22, 0), (22, 20)),
        "c": ((22, 20), (22, 40)), "d": ((0, 40), (22, 40)),
        "g": ((0, 20), (22, 20)),
    }
    for x, digits in ((510, "abc"), (548, "abgcd")):
        for segment in digits:
            p, q = segments[segment]
            draw.line((x + p[0], 297 + p[1], x + q[0], 297 + q[1]), fill="white", width=5)
    buffer = BytesIO()
    image.save(buffer, format="PNG", optimize=False, compress_level=9)
    return buffer.getvalue()


def _docx(image: bytes) -> bytes:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>"""
    relationships = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>"""
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">
<w:body><w:p><w:r><w:t>Synthetic vehicle reference</w:t></w:r></w:p>
<w:p><w:r><w:t>Inspect the illustration to answer questions about the vehicle.</w:t></w:r></w:p>
<w:p><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">
<wp:extent cx="5486400" cy="2926080"/><wp:docPr id="1" name="Vehicle illustration" descr="Original vehicle reference illustration"/>
<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
<pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="car-random-image.png"/><pic:cNvPicPr/></pic:nvPicPr>
<pic:blipFill><a:blip r:embed="rIdImage"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="5486400" cy="2926080"/></a:xfrm>
<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>
</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>
<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" w:gutter="0"/></w:sectPr>
</w:body></w:document>"""
    document_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/car-random-image.png"/>
</Relationships>"""
    core = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
 xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/"
 xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<dc:title>Synthetic vehicle reference</dc:title><dc:creator>Synthetic sales fixture</dc:creator>
<dcterms:created xsi:type="dcterms:W3CDTF">{FIXED_TIMESTAMP}</dcterms:created>
<dcterms:modified xsi:type="dcterms:W3CDTF">{FIXED_TIMESTAMP}</dcterms:modified>
</cp:coreProperties>"""
    app = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">
<Application>Synthetic sales fixture</Application></Properties>"""
    entries = {name: value.encode("utf-8") for name, value in {
        "[Content_Types].xml": content_types, "_rels/.rels": relationships,
        "word/document.xml": document, "word/_rels/document.xml.rels": document_rels,
        "docProps/core.xml": core, "docProps/app.xml": app,
    }.items()}
    entries["word/media/car-random-image.png"] = image
    return _stable_zip(entries)


def _oracle(partitions: dict[str, list[dict]]) -> dict:
    """Independent Decimal implementation: deliberately does not use runtime arithmetic."""
    observed = {}
    subtotals = {source: Decimal(0) for source in sorted(partitions)}
    source_details = {}
    with localcontext() as context:
        context.prec = 80
        for source in subtotals:
            rows = partitions[source]
            for row in rows:
                key = (row["dataset"], row["version"], row["sale_id"])
                if key in observed:
                    if observed[key] != row:
                        raise ValueError("oracle encountered a conflicting duplicate")
                    continue
                observed[key] = row
                if row["status"] != "cancelled":
                    subtotals[source] += (
                        Decimal(str(row["quantity"])) * Decimal(str(row["unit_price_minor"]))
                        - Decimal(str(row["discount_minor"])) - Decimal(str(row["refund_minor"]))
                    )
            ids = sorted(row["sale_id"] for row in rows)
            source_details[source] = {
                "row_count": len(rows), "sale_ids": ids,
                "sale_ids_sha256": _sha(("\n".join(ids) + "\n").encode("utf-8")),
                "rows_sha256": _sha(_rows_text(rows).encode("utf-8")),
            }
        total = sum(subtotals.values(), Decimal(0))
    return {
        "currency": "USD", "total_minor": int(total),
        "source_subtotals_minor": {source: int(amount) for source, amount in subtotals.items()},
        "input_row_count": sum(len(rows) for rows in partitions.values()),
        "unique_row_count": len(observed), "partitions": source_details,
        "oracle": "independent Decimal arithmetic, net pretax, cancelled=0; tax excluded",
        "overlap_attribution": "lexicographically first partition ID",
    }


def _rows_text(rows: list[dict]) -> str:
    return "".join(json.dumps(row, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
                   + "\n" for row in rows)


def generate(output: Path, version: str = "v1", seed: int = 48) -> dict:
    """Generate two distinct whole tables and return the answer-free publication.

    Output must be empty so stale artifacts cannot accidentally become ingested.
    ``publication.json`` describes bytes relative to either identical tree.
    Evaluation answers and intentionally overlapping fixtures live outside both.
    """
    if not isinstance(version, str) or not IDENTIFIER.fullmatch(version):
        raise ValueError("version must be a safe nonempty identifier")
    if type(seed) is not int:
        raise ValueError("seed must be an integer, not bool or float")
    output = Path(output)
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        raise ValueError("output directory must be empty; choose a new publication directory")
    encoding = tiktoken.encoding_for_model("text-embedding-3-large")
    if encoding.name != "cl100k_base":
        raise ValueError("expected cl100k_base tokenizer")
    partitions = {source: _partition_rows(source, version, seed, encoding) for source in ("json", "excel")}
    master = output / "master"
    for library in LIBRARY_LAYOUT:
        (master / library).mkdir(parents=True, exist_ok=True)
    saved_rows = _save_workbook(master / "originals" / "car-calculator-tool.xlsx", partitions["excel"])
    inputs = {"json": _json_table(partitions["json"], version), "excel": _worksheet_text(saved_rows)}
    publication_parts = []
    for source, rows in partitions.items():
        text = inputs[source]
        input_bytes = text.encode("utf-8")
        tokens = len(encoding.encode(text, disallowed_special=()))
        if not 13500 <= tokens <= 14500:
            raise ValueError("final whole-file token count is outside 13500..14500")
        input_path, rows_path = INPUT_PATHS[source], f"structured/{source}/rows.jsonl"
        row_text = _rows_text(rows)
        _write_text(master / input_path, text)
        _write_text(master / rows_path, row_text)
        ids = sorted(row["sale_id"] for row in rows)
        publication_parts.append({
            "id": source, "source": source, "site": "retail", "library": f"sales-{source}",
            "input_path": input_path, "rows_path": rows_path,
            "input_sha256": _sha(input_bytes), "rows_sha256": _sha(row_text.encode("utf-8")),
            "input_tokens": tokens, "row_count": len(rows),
            "sale_ids_sha256": _sha(("\n".join(ids) + "\n").encode("utf-8")),
            "expected_chunks": 3,
        })
    image = _draw_car()
    images = master / "images"
    images.mkdir()
    (images / "car-random-image.png").write_bytes(image)
    (images / "car-random-image.docx").write_bytes(_docx(image))
    files = [{"path": path.relative_to(master).as_posix(), "sha256": _sha(path.read_bytes())}
             for path in sorted(master.rglob("*")) if path.is_file()]
    publication = {
        "dataset": DATASET, "version": version, "seed": seed, "state": "generated",
        "partitions": publication_parts, "files": files,
    }
    expected = {
        "dataset": DATASET, "version": version, "seed": seed, **_oracle(partitions),
        "files": files, "image_expected_fact": IMAGE_FACT,
    }
    _write_text(output / "evaluation" / "expected.json", _json(expected))
    fixture_rng = random.Random(f"{DATASET}:{version}:{seed}:overlap-fixture")
    fixture_rows = [_sale(fixture_rng, number, version, fixture=True) for number in range(1, 5)]
    overlap = {"json": fixture_rows[:3], "excel": fixture_rows[1:]}
    _write_text(output / "fixtures" / "overlap.json", _json({
        "fixture_only": True, "not_for_ingestion": True,
        "partition_rows": overlap, "expected": _oracle(overlap),
    }))
    shutil.copytree(master, output / "project")
    for item in files:
        if (master / item["path"]).read_bytes() != (output / "project" / item["path"]).read_bytes():
            raise ValueError("master/project copies are not byte-identical")
    _write_text(output / "publication.json", _json(publication))
    return publication

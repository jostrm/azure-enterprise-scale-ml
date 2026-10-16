"""Normalize the two approved Graph sales sources into immutable local candidates.

Original business sale IDs and source versions are validated before rewriting the
row version to the content-derived publication version. Permission fingerprints
are separate lineage metadata so ACL-only synchronization can retain the same
input bytes, version and indexes without re-embedding. The exact
source workbook is retained separately from its normalized whole-table text.
This module neither reads nor writes an evaluation oracle or invokes cloud APIs.
"""
from __future__ import annotations

from io import BytesIO
import hashlib
import json
from pathlib import Path
import shutil
from zipfile import BadZipFile, ZipFile

import openpyxl
import tiktoken

from .arithmetic import validate_row
from .dataset import DATASET, FIELDS, INPUT_PATHS, LIBRARY_LAYOUT, _json_table, _rows_text, _worksheet_text
from .graph_delta import (
    GraphSyncError, IncompleteInventory, compact_item,
    project_permissions, validate_graph_config,
)


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _encoded(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n").encode()


def _validated_snapshots(config, snapshots):
    config = validate_graph_config(config)
    if not isinstance(snapshots, dict) or set(snapshots) != {"json", "excel"}:
        raise IncompleteInventory("Both required Graph partitions are necessary for a complete publication.")
    result = {}
    for binding in config["bindings"]:
        source = binding["partition"]
        snapshot = snapshots[source]
        if not isinstance(snapshot, dict):
            raise GraphSyncError("Malformed Graph source snapshot.")
        item = compact_item(snapshot.get("item"), config["drive_id"])
        if (item["id"] != binding["item_id"] or "deleted" in item or "file" not in item
                or set(item) & {"folder", "remoteItem", "package", "malware"} or not item.get("eTag")):
            raise IncompleteInventory("A required source partition is deleted, foreign or not a versioned file.")
        raw = snapshot.get("raw")
        if not isinstance(raw, bytes) or not raw or len(raw) > config["max_content_bytes"]:
            raise GraphSyncError("Source content must be nonempty bytes within the 16 MiB limit.")
        policy = project_permissions(config, snapshot.get("permissions"))
        result[source] = {"raw": raw, "item": item, "policy": policy}
    return config, result


def _version(config, snapshots):
    source_pins = {
        source: {"item_id": value["item"]["id"], "etag": value["item"]["eTag"],
                 "source_sha256": _sha(value["raw"])}
        for source, value in snapshots.items()}
    material = {"normalization": "graph-car-sales-v1", "tenant_id": config["tenant_id"],
                "drive_id": config["drive_id"], "scope_id": config["scope_id"], "sources": source_pins}
    return "graph-" + _sha(_encoded(material))[:24]


def candidate_version(config, snapshots):
    """Return graph-<24 hex digits> pinning bytes and source identity/ETag, not ACLs."""
    config, snapshots = _validated_snapshots(config, snapshots)
    return _version(config, snapshots)


def _json_rows(raw):
    def object_pairs(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise GraphSyncError("Duplicate JSON object keys are not allowed in source rows.")
            result[key] = value
        return result
    try:
        document = json.loads(raw.decode("utf-8-sig"), object_pairs_hook=object_pairs)
    except (UnicodeError, ValueError):
        raise GraphSyncError("The selected JSON source is not a valid whole sales table.") from None
    if (not isinstance(document, dict)
            or set(document) != {"dataset", "version", "currency", "records"}
            or document["dataset"] != DATASET or document["currency"] != "USD"
            or not isinstance(document["records"], list)):
        raise GraphSyncError("The selected JSON source must use the complete car-sales table schema.")
    return document["records"], document["version"]


def _excel_rows(raw):
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (len(entries) > 512 or len({entry.filename for entry in entries}) != len(entries)
                    or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024
                    or any("externallinks/" in entry.filename.lower()
                           or entry.filename.lower().endswith("vbaproject.bin") for entry in entries)):
                raise GraphSyncError("Workbook expansion, duplicate members or active external content is unsupported.")
        workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if workbook.sheetnames != ["Car sales"]:
                raise GraphSyncError("The approved workbook must contain exactly the Car sales worksheet.")
            sheet = workbook["Car sales"]
            if ((sheet.max_row and sheet.max_row > 20001)
                    or (sheet.max_column and sheet.max_column != len(FIELDS))):
                raise GraphSyncError("Workbook dimensions exceed the approved sales schema.")
            iterator = sheet.iter_rows()
            header = tuple(cell.value for cell in next(iterator))
            if header != FIELDS:
                raise GraphSyncError("Workbook header does not match the authoritative car-sales schema.")
            rows = []
            for cells in iterator:
                if len(rows) >= 20000 or any(cell.data_type == "f" for cell in cells):
                    raise GraphSyncError("Formula cells and oversized worksheets are not approved source data.")
                rows.append(dict(zip(FIELDS, (cell.value for cell in cells))))
            return rows, rows[0]["version"] if rows else None
        finally:
            workbook.close()
    except GraphSyncError:
        raise
    except (BadZipFile, ValueError, KeyError, StopIteration, TypeError, OSError):
        raise GraphSyncError("The selected Excel source is not a valid complete sales workbook.") from None


def _validate_rows(rows, version):
    if not rows or len(rows) > 20000:
        raise GraphSyncError("Each source requires a nonempty bounded complete row set.")
    seen = set()
    for row in rows:
        if not isinstance(row, dict) or set(row) != set(FIELDS):
            raise GraphSyncError("Every source row must match the complete car-sales schema.")
        validate_row(row, DATASET, version)
        if row["sale_id"] in seen:
            raise GraphSyncError("Duplicate business sale IDs within a source partition are forbidden.")
        seen.add(row["sale_id"])


def build_candidate(config, snapshots, destination, version=None):
    """Return an answer-free loader-compatible publication and write identical trees.

    No files are written until both complete sources, permissions, source version,
    arithmetic schema, overlaps and whole-file token budgets have been validated.
    If supplied, version must equal candidate_version(config, snapshots).
    Reader policies are attached to each partition for the caller to copy into its
    separately approved partition_policies; the runtime loader still owns that gate.
    """
    destination = Path(destination)
    if destination.exists() and (not destination.is_dir() or any(destination.iterdir())):
        raise GraphSyncError("Candidate output directory must be empty; immutable publications cannot be replaced.")
    if destination.is_symlink():
        raise GraphSyncError("Candidate output cannot be a symlink.")
    config, snapshots = _validated_snapshots(config, snapshots)
    derived = _version(config, snapshots)
    if version is not None and version != derived:
        raise GraphSyncError("The candidate version must match its content-derived version.")
    version = derived
    parsed = {"json": _json_rows(snapshots["json"]["raw"]),
              "excel": _excel_rows(snapshots["excel"]["raw"])}
    for source, (rows, source_version) in parsed.items():
        _validate_rows(rows, source_version)
    source_version = parsed["json"][1]
    if source_version != parsed["excel"][1]:
        raise GraphSyncError("Both source partitions must have one consistent original source version.")
    rows = {source: [{**row, "version": version} for row in values[0]]
            for source, values in parsed.items()}
    observed = {}
    for source in sorted(rows):
        for row in rows[source]:
            if row["sale_id"] in observed and row != observed[row["sale_id"]]:
                raise GraphSyncError("Conflicting cross-partition business sale IDs are forbidden.")
            observed[row["sale_id"]] = row
    encoding = tiktoken.get_encoding("cl100k_base")
    inputs = {"json": _json_table(rows["json"], version), "excel": _worksheet_text(rows["excel"])}
    files, partitions = {}, []
    for source in ("json", "excel"):
        raw = inputs[source].encode("utf-8")
        tokens = len(encoding.encode(inputs[source], disallowed_special=()))
        if not 9000 < tokens <= 18000:
            raise GraphSyncError("Whole inputs must exceed 9000 tokens and fit the three-chunk validation budget.")
        row_raw = _rows_text(rows[source]).encode("utf-8")
        input_path, rows_path = INPUT_PATHS[source], f"structured/{source}/rows.jsonl"
        files[input_path], files[rows_path] = raw, row_raw
        snapshot = snapshots[source]
        policy = snapshot["policy"]
        lineage = {"tenant_id": config["tenant_id"], "drive_id": config["drive_id"],
                   "item_id": snapshot["item"]["id"], "source_etag": snapshot["item"]["eTag"],
                   "source_sha256": _sha(snapshot["raw"]), "permission_sha256": policy["permission_sha256"],
                   "source_version": source_version, "normalization": "row-version-to-candidate-v1"}
        ids = sorted(row["sale_id"] for row in rows[source])
        partitions.append({
            "id": source, "source": source, "site": "retail", "library": f"sales-{source}",
            "input_path": input_path, "rows_path": rows_path, "input_sha256": _sha(raw),
            "rows_sha256": _sha(row_raw), "input_tokens": tokens, "row_count": len(rows[source]),
            "sale_ids_sha256": _sha(("\n".join(ids) + "\n").encode("utf-8")), "expected_chunks": 3,
            "graph_lineage": lineage, "allowed_user_ids": policy["allowed_user_ids"],
            "allowed_group_ids": policy["allowed_group_ids"],
        })
    files["originals/car-calculator-tool.xlsx"] = snapshots["excel"]["raw"]
    publication = {"dataset": DATASET, "version": version, "seed": 0, "state": "generated",
                   "partitions": partitions,
                   "files": [{"path": path, "sha256": _sha(raw)} for path, raw in sorted(files.items())]}
    created = False
    try:
        if not destination.exists():
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.mkdir()
            created = True
        for tree in ("master", "project"):
            root = destination / tree
            for library in LIBRARY_LAYOUT:
                root.joinpath(*library.split("/")).mkdir(parents=True, exist_ok=True)
            for path, raw in files.items():
                target = root.joinpath(*path.split("/"))
                target.parent.mkdir(parents=True, exist_ok=True)
                with target.open("xb") as stream:
                    stream.write(raw)
        for path, raw in files.items():
            if any(destination.joinpath(tree, *path.split("/")).read_bytes() != raw
                   for tree in ("master", "project")):
                raise GraphSyncError("Candidate master/project byte verification failed.")
        with (destination / "publication.json").open("xb") as stream:
            stream.write(_encoded(publication))
    except (OSError, ValueError, TypeError):
        # Only the new task-owned candidate is removed; never touch an existing tree.
        if created:
            shutil.rmtree(destination)
        raise
    return publication
